import openpyxl

wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)

control_skus = ['1502', '1503', '1504', '1505', '1506', '842', '848', '843', '002', '003', '005', '700\\703\\706\\708\\711']

with open('verify_output.txt', 'w', encoding='utf-8') as f:
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
            sku = row[0].value
            if sku in ['1502', '1503', '1504', '1505', '1506', '842', '848', '843', '002', '003', '005', '700\\703\\706\\708\\711']:
                type_val = row[2].value
                purpose_val = row[3].value
                name = row[1].value
                
                f.write(f"\nSKU: {sku}\n")
                f.write(f"  Name: {name}\n")
                f.write(f"  Type: {type_val}\n")
                f.write(f"  Purpose: {purpose_val}\n")

wb.close()