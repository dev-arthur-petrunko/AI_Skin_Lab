import openpyxl

wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)

with open('desc_output.txt', 'w', encoding='utf-8') as f:
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        f.write(f"\n=== Sheet: {sheet_name} ===\n")
        
        for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 10), min_col=1, max_col=12, values_only=False):
            sku = row[0].value
            name = row[1].value
            desc = row[11].value if len(row) > 11 else None
            type_val = row[2].value
            purpose_val = row[3].value
            f.write(f"SKU: {sku}\n")
            f.write(f"  Name: {name}\n")
            f.write(f"  Type: {type_val}\n")
            f.write(f"  Purpose: {purpose_val}\n")
            f.write(f"  Desc: {str(desc)[:500] if desc else None}\n")
            f.write("\n")

wb.close()