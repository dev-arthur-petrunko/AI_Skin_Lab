import openpyxl

wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)
ws = wb['Anubis']

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
    sku = row[0].value
    name = row[1].value
    type_val = row[2].value
    purpose_val = row[3].value
    desc = row[11].value if len(row) > 11 else None
    
    type_empty = type_val is None or (isinstance(type_val, str) and (type_val.strip() == '' or type_val.strip().lower() == 'nan'))
    purpose_empty = purpose_val is None or (isinstance(purpose_val, str) and (purpose_val.strip() == '' or purpose_val.strip().lower() == 'nan'))
    
    if type_val is None or (isinstance(type_val, str) and type_val.strip() == ''):
        print(f"SKU: {sku}")
        print(f"  Name: {name}")
        print(f"  Type: {repr(type_val)} (empty: {type_empty})")
        print(f"  Purpose: {repr(purpose_val)}")
        desc_text = str(row[11].value)[:200] if row[11].value else 'None'
        print(f"  Desc: {desc_text[:200]}")
        print()

wb.close()