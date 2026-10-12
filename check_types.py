import openpyxl

wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)
ws = wb['Anubis']

count = 0
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
    sku = row[0].value
    name = row[1].value
    type_val = row[2].value
    purpose_val = row[3].value
    desc = row[11].value if len(row) > 11 else None
    
    type_empty = type_val is None or (isinstance(type_val, str) and (type_val.strip() == '' or type_val.strip().lower() == 'nan'))
    purpose_empty = purpose_val is None or (isinstance(purpose_val, str) and (purpose_val.strip() == '' or purpose_val.strip().lower() == 'nan'))
    
    if type_val is None or (isinstance(type_val, str) and type_val.strip() == ''):
        with open('empty_rows.txt', 'a', encoding='utf-8') as f:
            f.write(f'SKU: {sku}\n')
            f.write(f'  Name: {name}\n')
            f.write(f'  Type: {repr(type_val)} (empty: {type_empty})\n')
            f.write(f'  Purpose: {repr(purpose_val)} (empty: {purpose_empty})\n')
            desc_text = str(desc)[:200] if desc else 'None'
            f.write(f'  Desc: {desc_text[:200]}\n\n')
            count += 1

print(f'Total empty rows: {count}')