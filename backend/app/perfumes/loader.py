"""Loader for the isolated perfume catalog (data/perfume.xlsx or a Perfume sheet).

Read-only: the Excel file is never modified. Falls back to the "perfume" sheet
inside price.xlsx when data/perfume.xlsx is absent (owner keeps the sheet there).
"""
from __future__ import annotations

import csv
import logging
import os
import re
import threading
import time
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

logger = logging.getLogger("perfumes.loader")

def _project_root() -> Path:
    """Works both on host (<repo>/backend/app/perfumes) and in container (/app/app/perfumes)."""
    env = os.getenv("PERFUME_DATA_DIR")
    if env:
        return Path(env).parent
    here = Path(__file__).resolve()
    for candidate in (here.parents[2], here.parents[3] if len(here.parents) > 3 else None):
        if candidate and list((candidate / "data").glob("*.xlsx")):
            return candidate
    return here.parents[2]


PROJECT_ROOT = _project_root()
DATA_DIR = PROJECT_ROOT / "data"
PERFUME_XLSX = Path(os.getenv("PERFUME_XLSX", str(DATA_DIR / "perfume.xlsx")))
PRICE_XLSX = Path(os.getenv("PRICE_XLSX", str(DATA_DIR / "price.xlsx")))
PERFUME_SHEET = os.getenv("PERFUME_SHEET", "perfume").strip().lower()
LOW_STOCK = int(os.getenv("PERFUME_LOW_STOCK", "5"))

LOOKALIKES = str.maketrans({
    "А": "A", "В": "B", "Е": "E", "К": "K", "М": "M", "Н": "H",
    "О": "O", "Р": "P", "С": "C", "Т": "T", "Х": "X", "а": "a",
    "о": "o", "с": "c", "е": "e", "р": "p", "х": "x",
})

LABEL_RE = re.compile(
    r"((?:верхн\w*|в\s+ерхн\w*)\s+нот\w*|нот\w*\s+серц\w*|середн\w*\s+нот\w*|баз\w*\s+нот\w*|шлейф\w*\s+нот\w*|нот\w*\s+шлейф\w*)\s*:",
    re.I,
)
SUSPICIOUS_SKUS = {"777777", "0", "1", "test"}

# Root scan: free text before the first label often contains marketing sentences.
# Families are extracted only when a known family root appears (order defines facet order).
FAMILY_ROOTS: list[tuple[str, str]] = [
    ("квіт", "Квіткові"), ("фрукт", "Фруктові"), ("цитрус", "Цитрусові"),
    ("орієнт", "Орієнтальні"), ("схід", "Орієнтальні"), ("восточн", "Орієнтальні"),
    ("дерев", "Деревні"), ("мускус", "Мускусні"), ("муск", "Мускусні"),
    ("ваніл", "Ванільні"), ("пряно", "Пряні"), ("прян", "Пряні"),
    ("шипро", "Шипрові"), ("амбр", "Амброві"), ("зелен", "Зелені"),
    ("гурман", "Гурманські"), ("аква", "Акватичні"), ("океан", "Акватичні"),
    ("водн", "Водні"), ("озон", "Озонічні"), ("пудр", "Пудрові"),
    ("фужер", "Фужерні"), ("ароматичн", "Ароматичні"), ("бурштин", "Бурштинові"),
    ("кожан", "Кожані"), ("мехов", "Мехові"), ("молочн", "Молочні"),
]


def _families_from(text: str) -> list[str]:
    out: list[str] = []
    for root, canon in FAMILY_ROOTS:
        if canon not in out and re.search(root, text, re.I):
            out.append(canon)
    return out


def _canon_sku(s: str) -> str:
    return unicodedata.normalize("NFC", s).translate(LOOKALIKES).strip()


def num(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v).replace("\u00a0", "").replace(" ", ""))
    return float(s.replace(",", ".")) if s else None


def photo_url(cell: Any) -> str | None:
    if cell is None:
        return None
    link = getattr(cell, "hyperlink", None)
    target = getattr(link, "target", None) if link else None
    if target:
        return str(target).strip()
    m = re.match(r'=HYPERLINK\("([^"]+)"', str(cell.value or ""), re.I)
    return m.group(1) if m else None


def _label_kind(label: str) -> str:
    low = label.lower().replace("в ерхн", "верхн")  # tolerate the "в ерхні ноти" data typo
    if low.startswith("верх"):
        return "top"
    if "серц" in low or low.startswith("серед"):
        return "heart"
    return "base"  # базові / шлейфові


def _split_list(s: str) -> list[str]:
    return [x.strip(" .;").capitalize() for x in re.split(r"[;,]", s) if x.strip(" .;")]


def parse_notes(raw: str) -> dict[str, Any]:
    raw = re.sub(r"\s+", " ", (raw or "").replace("\u00a0", " ")).strip()
    parts = LABEL_RE.split(raw)
    head = re.sub(r"(?i)г\s*рупа\s+ароматів\s*:?", "", parts[0])
    out: dict[str, Any] = {
        "families": _families_from(head),
        "top": [], "heart": [], "base": [],
        "description": "", "raw": raw,
    }
    recognized = False
    for lab, txt in zip(parts[1::2], parts[2::2]):
        recognized = True
        out[_label_kind(lab)] += [
            n for n in _split_list(txt)
            if re.search(r"[A-Za-zА-ЯІЇЄҐа-яіїєґ]", n)
        ]
    if not recognized:
        out["description"] = raw
    for k in ("top", "heart", "base"):
        seen: list[str] = []
        for n in out[k]:
            if n.lower() not in {s.lower() for s in seen}:
                seen.append(n)
        out[k] = seen
    return out


def _gender_type(name: str) -> tuple[str | None, str | None, str]:
    low = name.lower()
    gender = None
    if re.search(r"жіноч|для жінок", low):
        gender = "Жіночі"
    elif re.search(r"чоловіч|для чоловік", low):
        gender = "Чоловічі"
    elif "унісекс" in low or "unisex" in low:
        gender = "Унісекс"
    ptype = None
    if "туалетна вода" in low or "edt" in low:
        ptype = "Туалетна вода"
    elif "парфумован" in low or "edp" in low:
        ptype = "Парфумована вода"
    elif "парфуми" in low or re.search(r"\bпарфюм", low):
        ptype = "Парфуми"
    display = re.sub(
        r"^(жіноча|чоловіча|унісекс)\s+(парфумована вода|туалетна вода|парфуми)\s+(для жінок|для чоловіків|унісекс)?\s*",
        "", name, flags=re.I,
    ).strip(" ,")
    if len(display) < 3:
        display = re.sub(
            r"^(парфумована вода|туалетна вода|парфуми)\s+(для жінок|для чоловіків)?\s*",
            "", name, flags=re.I,
        ).strip(" ,")
    if len(display) < 3:
        display = name
    return gender, ptype, display


def _volume_ml(v: Any) -> int | None:
    n = num(v)
    if n is None:
        return None
    return int(round(n))


def _slug_base(sku: str) -> str:
    s = re.sub(r"[^0-9A-Za-zА-Яа-я]+", "-", _canon_sku(sku)).strip("-").lower()
    return s or "item"


def _resolve_source() -> tuple[Path | None, str | None, str | None]:
    """Return (path, sheet, reason)."""
    if PERFUME_XLSX.exists():
        return PERFUME_XLSX, None, None
    if PRICE_XLSX.exists():
        try:
            xl = load_workbook(PRICE_XLSX, read_only=True)
            names = xl.sheetnames
            xl.close()
        except Exception:  # noqa: BLE001
            names = []
        for n in names:
            if n.strip().lower() == PERFUME_SHEET:
                return PRICE_XLSX, n, None
        return None, None, f"немає data/perfume.xlsx і немає листа '{PERFUME_SHEET}' у price.xlsx"
    return None, None, "файл духів не знайдено"


def _cell(ws, r: int, col: dict[str, int | None], key: str):
    c = col.get(key)
    return ws.cell(r, c).value if c else None


def load_products() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Return (products, report_rows). Never raises for missing/broken file."""
    path, sheet, reason = _resolve_source()
    if path is None:
        logger.info("Perfume source unavailable: %s", reason)
        return [], [{"type": "source", "detail": reason or "unknown"}]

    try:
        wb = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:  # noqa: BLE001
        logger.exception("Cannot open %s", path)
        return [], [{"type": "source", "detail": str(exc)}]

    ws = wb[sheet] if sheet else None
    if ws is None:
        for n in wb.sheetnames:
            if n.strip().lower() == PERFUME_SHEET:
                ws = wb[n]
                break
    if ws is None:
        wb.close()
        return [], [{"type": "source", "detail": "лист не знайдено"}]

    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col: dict[str, int | None] = {}
    aliases = {
        "sku": ["Артикул", "SKU", "ID"],
        "name": ["Назва товару", "Назва", "Name"],
        "brand": ["Бренд", "Brand"],
        "price": ["Ціна, грн", "Ціна", "Price"],
        "promo": ["Промо ціна, грн", "Промо ціна", "Promo"],
        "discount": ["Знижка, %", "Знижка"],
        "stock": ["Залишок, шт", "Залишок", "Stock"],
        "volume": ["Об'єм", "Обʼєм", "Volume"],
        "country": ["Країна-виробник", "Країна", "Country"],
        "notes": ["Короткий опис (ноти)", "Опис", "Ноти"],
        "photo": ["Фото", "Image"],
    }
    for key, names in aliases.items():
        col[key] = next((i + 1 for i, h in enumerate(headers) if h in names), None)
    if not col["sku"] or not col["name"]:
        wb.close()
        return [], [{"type": "source", "detail": "немає колонок Артикул/Назва"}]

    products: list[dict[str, Any]] = []
    report: list[dict[str, Any]] = []
    slug_counts: dict[str, int] = {}
    seen_skus: dict[str, int] = {}

    for r in range(2, ws.max_row + 1):
        raw_sku = _cell(ws, r, col, "sku")
        if raw_sku is None or str(raw_sku).strip() == "":
            continue
        if isinstance(raw_sku, float) and raw_sku == int(raw_sku):
            raw_sku = int(raw_sku)
        sku = _canon_sku(str(raw_sku).strip())
        name = str(_cell(ws, r, col, "name") or "").strip()
        if not name:
            continue
        seen_skus[sku] = seen_skus.get(sku, 0) + 1
        slug = _slug_base(sku)
        slug_counts[slug] = slug_counts.get(slug, 0) + 1
        if slug_counts[slug] > 1:
            slug = f"{slug}-{slug_counts[slug]}"
        if sku in SUSPICIOUS_SKUS or len(sku) < 3:
            report.append({"type": "sku", "sku": sku, "name": name, "detail": "підозрілий/дубль артикул"})
        if seen_skus[sku] > 1:
            report.append({"type": "sku", "sku": sku, "name": name, "detail": f"дубль артикула (#{seen_skus[sku]})"})

        price = num(_cell(ws, r, col, "price")) or 0.0
        promo = num(_cell(ws, r, col, "promo"))
        col_discount = num(_cell(ws, r, col, "discount"))
        stock = int(num(_cell(ws, r, col, "stock")) or 0)
        discount = 0
        if promo is not None and price > 0 and promo < price:
            discount = round((1 - promo / price) * 100)
            if col_discount is not None and abs(col_discount - discount) > 1:
                report.append({
                    "type": "discount", "sku": sku, "name": name,
                    "detail": f"у файлі {col_discount:.0f}% розраховано {discount}%",
                })
        else:
            promo = None

        gender, ptype, display = _gender_type(name)
        parsed = parse_notes(str(_cell(ws, r, col, "notes") or ""))
        if parsed["raw"] and not parsed["top"] and not parsed["heart"] and not parsed["base"] and not parsed["description"]:
            report.append({"type": "notes", "sku": sku, "name": name, "detail": "нерозпізнаний опис нот"})

        url = photo_url(ws.cell(r, col["photo"]) if col["photo"] else None)
        volume = _volume_ml(_cell(ws, r, col, "volume"))
        low_stock = stock if 0 < stock <= LOW_STOCK else None
        products.append({
            "slug": slug,
            "sku": sku,
            "name": name,
            "display_name": display,
            "brand": str(_cell(ws, r, col, "brand") or "").strip(),
            "gender": gender,
            "type": ptype,
            "price": price,
            "promo_price": promo,
            "discount": discount,
            "stock": stock,
            "in_stock": stock > 0,
            "low_stock": low_stock,
            "volume_ml": volume,
            "country": str(_cell(ws, r, col, "country") or "").strip(),
            "families": parsed["families"],
            "notes": {"top": parsed["top"], "heart": parsed["heart"], "base": parsed["base"]},
            "note_list": (parsed["top"] + parsed["heart"] + parsed["base"]),
            "description": parsed["description"],
            "has_photo": False,
            "photo_url": url,
        })
    wb.close()
    report.append({"type": "summary", "detail": f"товарів: {len(products)}, унікальних артикулів: {len(seen_skus)}"})
    return products, report


def write_reports(products: list[dict[str, Any]], report: list[dict[str, Any]], photos: list[dict[str, Any]]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    with (DATA_DIR / "perfume_report.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["type", "sku", "name", "detail"], extrasaction="ignore")
        w.writeheader()
        w.writerows(report)
    with (DATA_DIR / "perfume_photos_report.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["sku", "slug", "url", "status"], extrasaction="ignore")
        w.writeheader()
        w.writerows(photos)


class PerfumeStore:
    """Thread-safe in-memory store with mtime-based auto-reload."""

    def __init__(self) -> None:
        self._lock = threading.Lock()
        self.products: list[dict[str, Any]] = []
        self.reason: str | None = None
        self.last_updated: str | None = None
        self._mtime: float | None = None
        self._callbacks: list = []

    def on_reload(self, cb) -> None:
        self._callbacks.append(cb)

    def _source_mtime(self) -> float | None:
        path, sheet, _ = _resolve_source()
        if path is None:
            return None
        try:
            m = path.stat().st_mtime
            if sheet:
                return m  # sheet edits bump the workbook mtime too
            return m
        except OSError:
            return None

    def reload(self) -> None:
        with self._lock:
            products, report = load_products()
            from .images import mark_products_photo  # local import to avoid cycle
            mark_products_photo(products)
            self.products = products
            self.reason = None if products else (report[0]["detail"] if report else "порожній файл")
            self._mtime = self._source_mtime()
            from datetime import datetime, timezone
            self.last_updated = datetime.now(timezone.utc).isoformat()
        for cb in self._callbacks:
            try:
                cb(products)
            except Exception:  # noqa: BLE001
                logger.exception("perfume reload callback failed")

    def ensure_fresh(self) -> None:
        m = self._source_mtime()
        if m != self._mtime:
            self.reload()

    def snapshot(self) -> tuple[list[dict[str, Any]], str | None]:
        self.ensure_fresh()
        return list(self.products), self.reason


store = PerfumeStore()
