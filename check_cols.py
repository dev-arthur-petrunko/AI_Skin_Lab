import openpyxl

wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)
print("Sheets:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=== Sheet: {sheet_name} ===")
    print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
    
    # Get headers from first row
    headers = []
    for cell in ws[1]:
        val = cell.value
        print(f"  Col {cell.column}: '{val}' (type: {type(val).__name__})")
    
    # Check if Призначення and Тип шкіри exist
    for cell in ws[1]:
        if cell.value:
            val = str(cell.value).strip().lower()
            if 'признач' in val:
                print(f"  Found Призначення at col {cell.column}: '{cell.value}'")
            elif 'тип' in val and ('шкір' in val or 'кожи' in val):
                print(f"  Found Тип шкіри at col {cell.column}: '{cell.value}'")

wb.close()