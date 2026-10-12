import openpyxl

wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)
ws = wb['Anubis']

count = 0
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
    sku = row[0].value
    name = row[1].value
    type_val = row[2].value
    purpose_val = row[3].value
    
    if sku and str(sku).strip() and str(sku).strip().lower() != 'nan':
        type_val = row[2].value
        purpose_val = row[3].value
        type_empty = type_val is None or (isinstance(type_val, str) and type_val.strip() == '')
        purpose_empty = purpose_val is None or (isinstance(purpose_val, str) and purpose_val.strip() == '')
        if type_empty or purpose_empty:
            print(f"SKU: {row[0].value} | Name: {str(name)[:50]} | Type empty: {type_empty} | Purpose empty: {purpose_empty}")
            count += 1

print(f"Empty rows: {count}")