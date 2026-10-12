import openpyxl

wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)
ws = wb['Anubis']

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
    sku = row[0].value
    if sku and '16060300' in str(sku):
        print(f'SKU: {repr(row[0].value)}')
        print(f'  Type: {repr(row[2].value)}')
        purpose_val = row[3].value
        print(f'  Purpose: {repr(row[3].value)[:100]}')
        type_val = row[2].value
        purpose_val = row[3].value
        if type_val:
            print(f'  Type bytes: {type_val.encode("utf-8")[:50]}')
        if row[3].value:
            print(f'  Purpose bytes: {row[3].value.encode("utf-8")[:100]}')
        break