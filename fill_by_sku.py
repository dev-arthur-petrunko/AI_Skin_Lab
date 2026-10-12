import openpyxl
import shutil
from datetime import datetime

# SKU-based mapping (more reliable than names)
SKU_MAPPING = {
    # Aqua O2XY line
    '1506': ('Зволоження; Антивіковий догляд; Відновлення', 'Нормальна; Суха; Чутлива'),  # Aqua O2XY Pro-Collagen Eye Cream
    '1505': ('Зволоження; Антивіковий догляд; Відновлення', 'Нормальна; Суха; Чутлива'),  # Aqua O2XY Pro-Collagen Night
    '1502': ('Зволоження; Відновлення; Заспокоєння та почервоніння; Освіження та тонізація', 'Нормальна; Суха; Чутлива'),  # Aqua O2XY Vitality Fresh Toner
    '1503': ('Зволоження; Антивіковий догляд; Відновлення', 'Нормальна; Суха; Чутлива'),  # Aqua O2XY Pro-Collagen Serum
    '1504': ('Зволоження; Антивіковий догляд; Сонцезахист', 'Нормальна; Суха; Чутлива'),  # Aqua O2XY Pro-Collagen Day SPF 30
    '1501': ('Очищення; Зволоження', 'Нормальна; Суха; Чутлива'),  # Aqua O2XY Velvet Glow Mousse
    '1503': ('Зволоження; Антивіковий догляд; Відновлення', 'Нормальна; Суха; Чутлива'),  # Aqua O2XY Pro-Collagen Serum 30
    '1504': ('Зволоження; Антивіковий догляд; Сонцезахист', 'Нормальна; Суха; Чутлива'),  # Aqua O2XY Pro-Collagen Day SPF 30
    
    # Illumination Quadro
    '842': ('Освітлення та пігментація; Антивіковий догляд; Зволоження; Антиоксидантний захист', 'Всі типи'),  # Illumination Quadro C Elixir
    
    # Smart4Derma Illumination
    '848': ('Освітлення та пігментація; Зволоження; Делікатні зони', 'Чутлива'),  # Illumination Toning Alginate Mask
    '843': ('Освітлення та пігментація; Зволоження; Делікатні зони', 'Чутлива'),  # Smart4Derma Illumination
    
    # Smart4Derma Acne Balance
    '003': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція; Заспокоєння та почервоніння', 'Жирна; Комбінована; Проблемна'),  # Acne Balance HomeKit
    '003': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція; Заспокоєння та почервоніння', 'Жирна; Комбінована; Проблемна'),
    
    # Smart4Derma Aquagen
    '005': ('Зволоження; Відновлення; Заспокоєння та почервоніння', 'Суха; Нормальна; Чутлива; Комбінована'),  # Aquagen HomeKit
    '005': ('Зволоження; Відновлення; Заспокоєння та почервоніння', 'Суха; Нормальна; Чутлива; Комбінована'),
    
    # Smart4Derma Age Performance
    '700\\703\\706\\708\\711': ('Антивіковий догляд; Зволоження; Освітлення та пігментація; Антиоксидантний захист', 'Всі типи'),  # Age Performance HomeKit
    '002': ('Антивіковий догляд; Зволоження; Освітлення та пігментація; Антиоксидантний захист', 'Всі типи'),  # Rediance Home Kit
    
    # CEF Lab B-Biotic
    '1606': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція; Очищення', 'Жирна; Комбінована; Проблемна'),  # Calamine Peel Mask
    '1605': ('Зволоження; Антивіковий догляд; Відновлення', 'Нормальна; Суха; Чутлива'),  # Balance Night Cream
    '1603': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція; Очищення', 'Жирна; Комбінована; Проблемна'),  # Sebo Active Serum
    '1604': ('Зволоження; Антивіковий догляд; Сонцезахист', 'Нормальна; Суха; Чутлива'),  # Balance Day SPF 30
    '1601': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція; Очищення', 'Жирна; Комбінована; Проблемна'),  # Sebo Active Cleanser
    '1602': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція; Очищення', 'Жирна; Комбінована; Проблемна'),  # Sebo Active Toner
    
    # CEF Lab B-Biotic other
    '1607': ('Зволоження; Антивіковий догляд; Сонцезахист', 'Нормальна; Суха; Чутлива'),  # Pro-Collagen Day SPF 30
    '1608': ('Антивіковий догляд; Зволоження; Відновлення', 'Нормальна; Суха; Чутлива'),  # Pro-Collagen Night
    '1609': ('Зволоження; Антивіковий догляд; Відновлення', 'Нормальна; Суха; Чутлива'),  # Pro-Collagen Serum
    '1610': ('Очищення; Зволоження', 'Нормальна; Суха; Чутлива'),  # Velvet Glow Mousse
    '1611': ('Антивіковий догляд; Зволоження; Відновлення', 'Нормальна; Суха; Чутлива'),  # Pro-Collagen Eye Cream
    
    # CEF Lab other
    '1201': ('Очищення; Зволоження; Заспокоєння та почервоніння', 'Чутлива; Нормальна; Суха'),  # Gentle Cleansing Gel
    '1202': ('Очищення; Зволоження; Заспокоєння та почервоніння', 'Чутлива; Нормальна; Суха'),  # Gentle Cleansing Gel 250
    '1203': ('Очищення; Себорегуляція; Матування та пори', 'Жирна; Комбінована; Проблемна'),  # Purifying Gel
    '1204': ('Освіження та тонізація; Себорегуляція; Матування та пори', 'Жирна; Комбінована; Проблемна'),  # Toner
    '1205': ('Зволоження; Антивіковий догляд; Сонцезахист', 'Нормальна; Суха; Чутлива'),  # Day Cream SPF 30
    '1206': ('Антивіковий догляд; Зволоження; Відновлення', 'Нормальна; Суха; Чутлива'),  # Night Cream
    
    # Anubis
    'A12031000': ('Очищення; Зволоження; Заспокоєння та почервоніння', 'Жирна; Комбінована; Проблемна; Чутлива'),  # Cleansing Gel 1000ml
    'A12030250': ('Очищення; Зволоження; Заспокоєння та почервоніння', 'Жирна; Комбінована; Проблемна; Чутлива'),  # Cleansing Gel 250ml
    'A2061000': ('Освіження та тонізація; Себорегуляція; Матування та пори', 'Жирна; Комбінована; Проблемна'),  # Tonifying Lotion 2%
    'A12060250': ('Освіження та тонізація; Себорегуляція; Матування та пори', 'Жирна; Комбінована; Проблемна'),  # Tonifying Lotion 2% 250ml
    'A12140050': ('Очищення; Освітлення та пігментація; Матування та пори', 'Жирна; Комбінована; Проблемна'),  # Glyco Gel 15/10
    'A12130250': ('Очищення; Зволоження; Сонцезахист', 'Нормальна; Суха; Комбінована'),  # Oxygen Cream SPF 15
    'A15080050': ('Освітлення та пігментація; Антивіковий догляд; Зволоження', 'Всі типи'),  # MELA TRX Whitening Emulsion
    'A15030010': ('Освітлення та пігментація; Антивіковий догляд; Зволоження', 'Всі типи'),  # Whitening Roll On
    'A12060250(1)': ('Освіження та тонізація; Себорегуляція; Матування та пори', 'Жирна; Комбінована; Проблемна'),  # Tonifying Lotion 2% 250ml (duplicate)
    'A12140050': ('Очищення; Освітлення та пігментація; Матування та пори', 'Жирна; Комбінована; Проблемна'),  # Glyco Gel 15/10 (duplicate)
    'A12130250': ('Очищення; Зволоження; Сонцезахист', 'Нормальна; Суха; Комбінована'),  # Oxygen Cream SPF 15 (duplicate)
    
    # Sets
    '002': ('Антивіковий догляд; Зволоження; Освітлення та пігментація; Антиоксидантний захист', 'Всі типи'),  # REDIANCE HOME KIT
    '003': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція; Заспокоєння та почервоніння', 'Жирна; Комбінована; Проблемна'),  # Acne Balance HomeKit
    '005': ('Зволоження; Відновлення; Заспокоєння та почервоніння', 'Суха; Нормальна; Чутлива; Комбінована'),  # Aquagen HomeKit
    '700\\703\\706\\708\\711': ('Антивіковий догляд; Зволоження; Освітлення та пігментація; Антиоксидантний захист', 'Всі типи'),  # Age Performance HomeKit
    '005': ('Зволоження; Відновлення; Заспокоєння та почервоніння', 'Суха; Нормальна; Чутлива; Комбінована'),  # Aquagen HomeKit (duplicate)
    '002': ('Антивіковий догляд; Зволоження; Освітлення та пігментація; Антиоксидантний захист', 'Всі типи'),  # Rediance Home Kit (duplicate)
}

# Brand-based fallbacks
BRAND_DEFAULTS = {
    'smart4derma': ('Зволоження; Антивіковий догляд; Освітлення та пігментація', 'Нормальна; Суха; Чутлива'),
    'cef lab': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція; Очищення', 'Жирна; Комбінована; Проблемна'),
    'cef lab b-biotic': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція; Очищення', 'Жирна; Комбінована; Проблемна'),
    'cef lab aqua': ('Зволоження; Антивіковий догляд; Відновлення', 'Нормальна; Суха; Чутлива'),
    'cef lab aqua o2xy': ('Зволоження; Антивіковий догляд; Відновлення', 'Нормальна; Суха; Чутлива'),
    'anna logon': ('Матування та пори; Акне та проблемна шкіра; Себорегуляція', 'Жирна; Комбінована'),
    'cef lab b-biotic': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція; Очищення', 'Жирна; Комбінована; Проблемна'),
    'anubis': ('Очищення; Себорегуляція; Матування та пори', 'Жирна; Комбінована; Проблемна'),
    'smart4derma': ('Зволоження; Антивіковий догляд; Освітлення та пігментація', 'Нормальна; Суха; Чутлива'),
    'smart4derma aqua': ('Зволоження; Відновлення; Заспокоєння та почервоніння', 'Суха; Нормальна; Чутлива; Комбінована'),
    'smart4derma acne': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція', 'Жирна; Комбінована; Проблемна'),
    'smart4derma age': ('Антивіковий догляд; Зволоження; Освітлення та пігментація; Антиоксидантний захист', 'Всі типи'),
    'anubis': ('Очищення; Себорегуляція; Матування та пори', 'Жирна; Комбінована; Проблемна'),
    'smart4derma rediance': ('Освітлення та пігментація; Антивіковий догляд; Зволоження', 'Нормальна; Суха; Комбінована'),
    'smart4derma aquagen': ('Зволоження; Відновлення; Заспокоєння та почервоніння', 'Суха; Нормальна; Чутлива; Комбінована'),
    'smart4derma acne': ('Акне та проблемна шкіра; Матування та пори; Себорегуляція', 'Жирна; Комбінована; Проблемна'),
    'smart4derma age': ('Антивіковий догляд; Зволоження; Освітлення та пігментація; Антиоксидантний захист', 'Всі типи'),
    'illumination': ('Освітлення та пігментація; Антивіковий догляд; Зволоження; Антиоксидантний захист', 'Всі типи'),
    'anubis': ('Очищення; Себорегуляція; Матування та пори', 'Жирна; Комбінована; Проблемна'),
}

GENERIC_PURPOSE = 'Зволоження; Очищення; Відновлення'
GENERIC_TYPE = 'Нормальна; Суха; Чутлива'

def find_match(sku, name='', brand=''):
    # Try exact SKU match first
    if sku in SKU_MAPPING:
        return SKU_MAPPING[sku]
    
    # Try brand-based fallback
    name_lower = str(sku).lower() + ' ' + str(name).lower()
    for brand_key, (purpose, skin_type) in BRAND_DEFAULTS.items():
        if brand_key in str(sku).lower() + ' ' + str(name).lower():
            return purpose, skin_type
    
    return None, None

# ============ MAIN ============
import openpyxl
import shutil
from datetime import datetime

print("Loading workbook...")
wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)

# Backup
backup_path = f'data/backup/price_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
shutil.copy2('data/price.xlsx', backup_path)
print(f"Backup created: {backup_path}")

total_filled_type = 0
total_filled_purpose = 0
report_rows = []

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\nProcessing sheet: {sheet_name}")
    
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False), start=2):
        sku_cell = row[0]
        name_cell = row[1]
        type_cell = row[2]
        purpose_cell = row[3]
        
        sku = sku_cell.value
        name = name_cell.value
        
        if not sku:
            continue
            
        type_val = type_cell.value
        purpose_val = purpose_cell.value
        
        type_empty = type_val is None or (isinstance(type_val, str) and type_val.strip() == '')
        purpose_empty = purpose_val is None or (isinstance(purpose_val, str) and purpose_val.strip() == '')
        
        if not type_empty and not purpose_empty:
            continue
        
        # Find match
        name_str = str(name_cell.value) if name_cell.value else ''
        purpose, skin_type = find_match(sku, name_cell.value, '')
        
        filled = False
        if purpose_empty and purpose:
            purpose_cell.value = purpose
            total_filled_purpose += 1
        
        if type_empty and skin_type:
            type_cell.value = skin_type
            total_filled_type += 1

# Save
wb.save('data/price.xlsx')
print(f"\nFilled Тип шкіри: {total_filled_type}")
print(f"Filled Призначення: {total_filled_purpose}")

# Verify controls
wb2 = openpyxl.load_workbook('data/price.xlsx', data_only=False)
control_skus = {
    'aqua o2xy': ['1502', '1503', '1504', '1505', '1506'],
    'illumination quadro': ['842'],
    'smart4derma illumination': ['848', '843'],
}

for sheet_name in wb2.sheetnames:
    ws = wb2[sheet_name]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
        sku = row[0].value
        name = row[1].value
        type_val = row[2].value
        purpose_val = row[3].value
        
        for control_name, skus in {
            'aqua o2xy': ['1502', '1503', '1504', '1505', '1506'],
            'illumination quadro': ['842'],
            'smart4derma illumination': ['848', '843'],
        }.items():
            if sku in skus:
                print(f"\n=== CONTROL: {control_name} (SKU: {sku}) ===")
                print(f"  Name: {row[1].value}")
                print(f"  Type: {type_val}")
                print(f"  Purpose: {purpose_val}")

wb.save('data/price.xlsx')
print(f"\nSaved.")