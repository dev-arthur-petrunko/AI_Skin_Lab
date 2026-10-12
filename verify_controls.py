import openpyxl

wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)

with open('control_check.txt', 'w', encoding='utf-8') as f:
    control_skus = {
        'Aqua O2XY': ['Aqua O2XY', 'Aqua O2XY'],
        'Illumination Quadro': ['Illumination Quadro', 'Quadro C Elixir'],
        'Smart4Derma Illumination': ['Smart4Derma Illumination', 'пінка для делікатних зон'],
        'CEF Lab Aqua O2XY Vitality Fresh Toner': ['CEF Lab Aqua O2XY Vitality Fresh Toner'],
        'Illumination Quadro C Elixir': ['Illumination Quadro C Elixir', 'Quadro C Elixir'],
        'Smart4Derma Illumination pinka': ['Smart4Derma Illumination', 'пінка для делікатних зон'],
    }

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
            name = row[1].value
            type_val = row[2].value
            purpose_val = row[3].value
            desc = row[11].value if len(row) > 11 else None
            
            name_str = str(name) if name else ''
            for control_name, keywords in control_skus.items():
                if any(kw.lower() in str(name).lower() for kw in keywords):
                    f.write(f"\n=== CONTROL: {control_name} ===\n")
                    f.write(f"  SKU: {row[0].value}\n")
                    f.write(f"  Name: {name}\n")
                    f.write(f"  Type: {type_val}\n")
                    f.write(f"  Purpose: {purpose_val}\n")
                    f.write(f"  Desc: {str(desc)[:300] if desc else None}\n")

wb.close()