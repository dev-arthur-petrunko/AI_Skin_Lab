import openpyxl

wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)

control_skus = ['1502', '1503', '1504', '1505', '1506', '842', '848', '843', '002', '003', '005', '700\\703\\706\\708\\711']

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
        sku = row[0].value
        if sku in control_skus:
            type_cell = row[2]
            purpose_cell = row[3]
            name = row[1].value
            
            type_val = type_cell.value
            purpose_val = purpose_cell.value
            
            print(f"\nSKU: {sku}")
            print(f"  Type: {repr(type_val)}")
            print(f"  Purpose: {repr(purpose_val)}")

wb.close()