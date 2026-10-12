import openpyxl

wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)

# Check what's in Type and Purpose columns for control SKUs
control_skus = ['1502', '1503', '1504', '1505', '1506', '842', '848', '843', '002', '003', '005', '700\\703\\706\\708\\711']

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
        sku = row[0].value
        if sku in control_skus:
            type_cell = row[2]
            purpose_cell = row[3]
            name = row[1].value
            
            type_val = type_cell.value
            purpose_val = purpose_cell.value
            
            print(f"\nSKU: {sku}")
            print(f"  Name: {row[1].value}")
            print(f"  Type cell value: {repr(type_cell.value)}")
            print(f"  Type empty: {type_cell.value is None or (isinstance(type_cell.value, str) and type_cell.value.strip() == '')}")
            print(f"  Purpose cell value: {repr(purpose_cell.value)}")
            print(f"  Purpose empty: {purpose_cell.value is None or (isinstance(purpose_cell.value, str) and purpose_cell.value.strip() == '')}")

wb.close()