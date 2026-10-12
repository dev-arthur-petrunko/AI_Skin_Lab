import openpyxl
import re
import shutil
from datetime import datetime
from pathlib import Path

# ============ DICTIONARIES ============
PURPOSE_KEYWORDS = {
    'зволож': ['зволож', 'hydrat', 'moistur', 'гиалурон', 'гидру', 'влажн'],
    'відновл': ['відновл', 'регенерац', 'репейр', 'repair', 'нориш', 'nourish'],
    'заспокоєнн': ['заспоко', 'спок', 'calm', 'sooth', 'купероз', 'розецеа', 'розацеа', 'почервон'],
    'освітленн': ['освітл', 'піднижен', 'тон', 'віт c', 'вітамін c', 'пigмент', 'спотил', 'bright', 'white', 'illumina', 'radiance', 'пigмент', 'pigment'],
    'антивіков': ['антивік', 'омолодж', 'ліфтинг', 'ретинол', 'ретінал', 'пепт', 'перепід', 'wrinkl', 'anti.age', 'anti-age', 'age'],
    'антиоксид': ['антиокс', 'антіокс', 'вітамін e', 'ферулова', 'coenzyme', 'q10', 'antiox'],
    'очищенн': ['очищ', 'пена', 'гель', 'мус', 'мусс', 'гідрофільн', 'зімне масло', 'cleans', 'wash', 'foam'],
    'тоніз': ['тоніз', 'тонік', 'лосьйон', 'toner', 'locion'],
    'матуванн': ['матув', 'себо', 'пор', 'пір', 'глян', 'blot', 'matt', 'oil.control'],
    'акне': ['акне', 'acne', 'проблем', 'вугр', 'pimple', 'blemish', 'sebo.active'],
    'сонцезах': ['spf', 'сонцезах', 'sunscreen', 'sun.protect', 'uv', 'уф'],
    'оч': ['оч', 'перелічк', 'eye', 'contour'],
    'делікатн': ['інтим', 'бікіні', 'делікат', 'intimate', 'bikini'],
}

SKIN_TYPE_KEYWORDS = {
    'сух': ['сух', 'dry', 'дегідрат'],
    'жирн': ['жирн', 'oily', 'себо', 'матув', 'глян', 'комбін', 't-zone'],
    'чутлив': ['чутлив', 'sensitive', 'сенсив', 'геперчут', 'атонік', 'купероз'],
    'нормальн': ['нормальн', 'normal'],
    'комбін': ['комбін', 'combined', 'mixed', 't-zone', 't-зона'],
    'акне': ['акне', 'acne', 'проблем', 'вугр', 'pimple'],
    'всі': ['всі типи', 'all skin', 'all skin types', 'універсальн'],
}

def normalize(text):
    if not text:
        return ''
    return str(text).lower().replace('’', "'").replace('–', '-').replace('—', '-')

def extract_purpose(text, name=''):
    if not text:
        text = ''
    text = normalize(text + ' ' + name)
    found = set()
    for key, kws in PURPOSE_KEYWORDS.items():
        for kw in kws:
            if kw in text:
                if key == 'зволож':
                    found.add('Зволоження')
                elif key == 'відновл':
                    found.add('Відновлення')
                elif key == 'заспокоєнн':
                    found.add('Заспокоєння та почервоніння')
                elif key == 'освітленн':
                    found.add('Освітлення та пігментація')
                elif key == 'антивіков':
                    found.add('Антивіковий догляд')
                elif key == 'антиоксид':
                    found.add('Антиоксидантний захист')
                elif key == 'очищенн':
                    found.add('Очищення')
                elif key == 'тоніз':
                    found.add('Освіження та тонізація')
                elif key == 'матуванн':
                    found.add('Матування та пори')
                elif key == 'акне':
                    found.add('Акне та проблемна шкіра')
                elif key == 'сонцезах':
                    found.add('Сонцезахист')
                elif key == 'оч':
                    found.add('Догляд за очима')
                elif key == 'делікатн':
                    found.add('Делікатні зони')
    return '; '.join(sorted(found)) if found else ''

def extract_skin_type(text, name=''):
    if not text:
        text = ''
    text = normalize(text + ' ' + name)
    found = set()
    for skin, kws in SKIN_TYPE_KEYWORDS.items():
        for kw in kws:
            if kw in text:
                if skin == 'сух':
                    found.add('Суха')
                elif skin == 'жирн':
                    found.add('Жирна')
                elif skin == 'чутлив':
                    found.add('Чутлива')
                elif skin == 'нормальн':
                    found.add('Нормальна')
                elif skin == 'комбін':
                    found.add('Комбінована')
                elif skin == 'акне':
                    pass  # don't add as skin type
                elif skin == 'всі':
                    found.add('Всі типи')
    return '; '.join(sorted(found)) if found else ''

def clean_value(val):
    if not val:
        return ''
    val = str(val).strip()
    parts = re.split(r'[;,\n•]', val)
    cleaned = []
    for p in parts:
        p = p.strip()
        p = re.sub(r'[^\w\s\-/+%().\']', '', p)
        p = p.strip(' .')
        if p:
            cleaned.append(p)
    return '; '.join(sorted(set(cleaned))) if cleaned else ''

# ============ MAIN ============
print("Loading workbook...")
wb = openpyxl.load_workbook('data/price.xlsx', data_only=False)

# Backup
backup_path = f'data/backup/price_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
shutil.copy2('data/price.xlsx', backup_path)
print(f"Backup created: {backup_path}")

total_filled_type = 0
total_filled_purpose = 0
report_rows = []
manual_fill = []

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\nProcessing sheet: {sheet_name}")
    
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False), start=2):
        sku_cell = row[0]
        name_cell = row[1]
        type_cell = row[2]
        purpose_cell = row[3]
        desc_cell = row[11] if len(row) > 11 else None
        
        sku = sku_cell.value
        name = name_cell.value
        desc = desc_cell.value if desc_cell else None
        
        if not sku:
            continue
            
        type_val = type_cell.value
        purpose_val = purpose_cell.value
        
        type_empty = type_val is None or (isinstance(type_val, str) and type_val.strip() == '')
        purpose_empty = purpose_val is None or (isinstance(purpose_val, str) and purpose_val.strip() == '')
        
        if not type_empty and not purpose_empty:
            continue
        
        name_text = str(name_cell.value) if name_cell.value else ''
        desc_text = str(desc_cell.value) if desc_cell and desc_cell.value else ''
        
        # Extract purpose and skin type
        combined = f"{name} {desc}".lower()
        
        # Purpose
        purpose = ''
        if purpose_empty:
            purpose = extract_purpose(desc_text + ' ' + name_text, name_text)
            if purpose:
                purpose_cell.value = purpose
                total_filled_purpose += 1
        
        # Skin type
        skin_type = ''
        if type_empty:
            skin_type = extract_skin_type(desc_text + ' ' + name_text, name_text)
            if skin_type:
                type_cell.value = skin_type
                total_filled_type += 1

# Save
wb.save('data/price.xlsx')
print(f"\nFilled Тип шкіри: {total_filled_type}")
print(f"Filled Призначення: {total_filled_purpose}")

# Verify controls
print("\n=== VERIFICATION ===")
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
        name = row[1].value
        type_val = row[2].value
        purpose_val = row[3].value
        desc = row[11].value if len(row) > 11 else None
        
        name_str = str(name) if name else ''
        if any(kw.lower() in str(name).lower() for kw in ['aqua o2xy', 'illumination quadro', 'smart4derma illumination']):
            print(f"\n=== CONTROL ===")
            print(f"  SKU: {row[0].value}")
            print(f"  Name: {name}")
            print(f"  Type: {type_val}")
            print(f"  Purpose: {purpose_val}")

wb.save('data/price.xlsx')
print(f"\nSaved.")